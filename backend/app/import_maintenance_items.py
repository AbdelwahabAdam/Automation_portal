"""Add the items طلب صيانة وتبديل offers (Maintenance_Invoice_Items.xlsx) to our product catalog.

Usage: python -m app.import_maintenance_items [path-to-Maintenance_Invoice_Items.xlsx]

Each row gives the item's group (التصنيف), its name on the invoice (also its Daftra product's name)
and its name in the portal. A product already in the catalog under that name (or Daftra id) is kept,
with its price, and only marked as a maintenance item; a new one starts at the price below. Prices,
names and which products are maintenance items are then managed in الإدارة ← المنتجات والباقات.
Safe to run again.
"""

import argparse
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select

from .database import SessionLocal, run_migrations
from .models import Product

# Each item's Daftra product: the one used on invoices where Daftra has several with the same name
DAFTRA_IDS = {
    "RO .PP 10": 625, "RO CTO 10": 627, "RO UDF 10": 626, "J PP 20": 628, "J CTO 20": 629,
    "RO Booster Pump Transformer / Adapter": 890, 'Bib Cock "1/2': 618, "fittings metal ball valve 1/4 inch": 896,
    "BRASS NIPPLE 1/2": 915, 'Float Valve 1/2" PVC': 920, "RO Membrane Elbow Check Valve": 921,
    "MULTIPORT VALVE F56K": 662, "GRAVEL SIZE 1-3": 635, "GRAVEL SIZE 3-5": 636, "SAND SIZE 04-08": 634,
    "Water Connection Service": 212, "service": 194,
}
# The price each was most often invoiced at on the last 1,200 invoices (checked 2026-10-21); a starting
# point only, changed in the admin page
START_PRICES = {
    "RO .PP 10": 50, "RO CTO 10": 30, "RO UDF 10": 40, "J PP 20": 50, "J CTO 20": 40,
    "RO Booster Pump Transformer / Adapter": 120, 'Bib Cock "1/2': 35, "fittings metal ball valve 1/4 inch": 35,
    "BRASS NIPPLE 1/2": 30, 'Float Valve 1/2" PVC': 100, "RO Membrane Elbow Check Valve": 150,
    "MULTIPORT VALVE F56K": 725, "GRAVEL SIZE 1-3": 150, "GRAVEL SIZE 3-5": 150, "SAND SIZE 04-08": 150,
    "Water Connection Service": 150, "service": 300,
}


def read_items(path: str) -> list[tuple[str, str, str]]:
    sheet = load_workbook(path, data_only=True).worksheets[0]
    rows = []
    for group, invoice_name, portal_name in sheet.iter_rows(min_row=5, max_col=3, values_only=True):
        if invoice_name:
            rows.append((str(group).strip(), " ".join(str(invoice_name).split()), str(portal_name or "").strip()))
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", nargs="?", default="../../Maintenance_Invoice_Items.xlsx")
    args = parser.parse_args()
    items = read_items(args.workbook)
    run_migrations()
    added, marked = [], []
    with SessionLocal() as db:
        for group, name, name_ar in items:
            daftra_id = DAFTRA_IDS.get(name)
            # The product with this name, else the one with its Daftra id
            product = db.scalar(select(Product).where(func.lower(Product.name) == name.lower()))
            if product is None and daftra_id:
                product = db.scalar(select(Product).where(Product.daftra_id == daftra_id))
            if product is None:
                product = Product(name=name, unit_price=Decimal(START_PRICES.get(name, 0)), price_includes_tax=True,
                                  tax_rate=Decimal("0.05"), is_active=True)
                db.add(product)
                added.append(name)
            else:
                marked.append(name)
            product.name_ar = product.name_ar or name_ar
            db.flush()
            # A Daftra id another product already carries stays with that product (it is unique)
            taken = daftra_id and db.scalar(select(Product.id).where(Product.daftra_id == daftra_id, Product.id != product.id))
            if not product.daftra_id and not taken:
                product.daftra_id = daftra_id
            product.maintenance_group = group
        db.commit()
    print(f"{len(items)} maintenance items: {len(added)} added to the catalog, {len(marked)} already there and marked")
    for name in added:
        print(f"  + {name} ({START_PRICES.get(name, 0)})")


if __name__ == "__main__":
    main()
